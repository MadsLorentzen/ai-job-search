from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from tests.webapp.api.test_review_routes import _app_with_pack_chain
from webapp.persistence.artifacts import get_current_artifact
from webapp.persistence.db import connect


def _confirm_pack(client, workspace_id):
    response = client.post(f"/api/workspaces/{workspace_id}/application-pack", json={
        "confirmed": True, "effective_date": "2026-08-24",
    })
    assert response.status_code == 201, response.text
    return response.json()


def test_download_cv_document_after_pack_confirmation(tmp_path):
    app, settings, workspace_id = _app_with_pack_chain(tmp_path)
    with TestClient(app) as client:
        _confirm_pack(client, workspace_id)
        response = client.get(f"/api/workspaces/{workspace_id}/application-pack/render/cv")
        assert response.status_code == 200
        assert response.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        assert "attachment" in response.headers["content-disposition"]
        assert response.headers["content-disposition"].endswith('.docx"')
        document = Document(BytesIO(response.content))
        texts = [p.text for p in document.paragraphs]
        assert any(text for text in texts)


def test_download_cover_letter_document_after_pack_confirmation(tmp_path):
    app, settings, workspace_id = _app_with_pack_chain(tmp_path)
    with TestClient(app) as client:
        _confirm_pack(client, workspace_id)
        response = client.get(
            f"/api/workspaces/{workspace_id}/application-pack/render/cover_letter"
        )
        assert response.status_code == 200
        document = Document(BytesIO(response.content))
        texts = [p.text for p in document.paragraphs]
        assert any(text for text in texts)


def test_render_before_pack_confirmation_is_a_clean_error(tmp_path):
    app, settings, workspace_id = _app_with_pack_chain(tmp_path)
    with TestClient(app) as client:
        response = client.get(f"/api/workspaces/{workspace_id}/application-pack/render/cv")
        assert response.status_code == 400


def test_render_unknown_kind_is_404(tmp_path):
    app, settings, workspace_id = _app_with_pack_chain(tmp_path)
    with TestClient(app) as client:
        _confirm_pack(client, workspace_id)
        response = client.get(
            f"/api/workspaces/{workspace_id}/application-pack/render/resume_video"
        )
        assert response.status_code == 404


def test_render_of_unknown_workspace_is_404(tmp_path):
    app, settings, workspace_id = _app_with_pack_chain(tmp_path)
    with TestClient(app) as client:
        response = client.get("/api/workspaces/does-not-exist/application-pack/render/cv")
        assert response.status_code == 404


def test_explicit_historical_pack_artifact_id_renders_that_exact_pack(tmp_path):
    app, settings, workspace_id = _app_with_pack_chain(tmp_path)
    with TestClient(app) as client:
        first_confirm = _confirm_pack(client, workspace_id)
        first_pack_artifact_id = first_confirm["artifact"]["id"]

        response = client.get(
            f"/api/workspaces/{workspace_id}/application-pack/render/cv",
            params={"pack_artifact_id": first_pack_artifact_id},
        )
        assert response.status_code == 200

    conn = connect(settings.db_path)
    current = get_current_artifact(conn, workspace_id, "application_pack")
    conn.close()
    assert current["id"] == first_pack_artifact_id


def test_download_remains_available_after_workflow_advances_past_drafted(tmp_path):
    """Regression: the immutable pack, and therefore its rendered documents,
    must stay downloadable once the workflow moves to applied/interview/etc.
    Downloading is not submission, and pack existence -- not workflow status
    -- is what should gate availability."""
    app, settings, workspace_id = _app_with_pack_chain(tmp_path)
    with TestClient(app) as client:
        _confirm_pack(client, workspace_id)

        applied = client.patch(f"/api/workspaces/{workspace_id}/status", json={
            "new_status": "applied", "effective_date": "2026-08-24",
        })
        assert applied.status_code == 200, applied.text

        response = client.get(f"/api/workspaces/{workspace_id}/application-pack/render/cv")
        assert response.status_code == 200

        interview = client.patch(f"/api/workspaces/{workspace_id}/status", json={
            "new_status": "interview", "effective_date": "2026-08-24",
        })
        assert interview.status_code == 200, interview.text

        response = client.get(
            f"/api/workspaces/{workspace_id}/application-pack/render/cover_letter"
        )
        assert response.status_code == 200


def test_render_rejects_pack_artifact_id_belonging_to_a_different_workspace(tmp_path):
    app, settings, workspace_id = _app_with_pack_chain(tmp_path)
    with TestClient(app) as client:
        first_confirm = _confirm_pack(client, workspace_id)
        first_pack_artifact_id = first_confirm["artifact"]["id"]

    other_app, other_settings, other_workspace_id = _app_with_pack_chain(tmp_path.parent / "other")
    with TestClient(other_app) as other_client:
        response = other_client.get(
            f"/api/workspaces/{other_workspace_id}/application-pack/render/cv",
            params={"pack_artifact_id": first_pack_artifact_id},
        )
        assert response.status_code == 400
