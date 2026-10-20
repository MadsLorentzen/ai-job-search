"""Run the converter on workbooks with both tabular and chart-only sheets."""
import json
from pathlib import Path
import subprocess
import sys
from tempfile import TemporaryDirectory
import unittest

try:
    import openpyxl
    from openpyxl.chart import BarChart, Reference
except ImportError:
    openpyxl = None


TOOL = Path(__file__).resolve().parent.parent / "tools" / "convert_salary_excel.py"


@unittest.skipUnless(openpyxl, "requires optional openpyxl (installed in CI)")
class ChartsheetCLI(unittest.TestCase):
    def convert(self, with_chart):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            workbook = root / "salary report.xlsx"
            output = root / "salary output.json"
            book = openpyxl.Workbook()
            try:
                first = book.active
                first.title = "Salaries"
                first.append(["Company", "Count", "Index"])
                first.append(["Acme Corp", 12, 105])
                second = book.create_sheet("Engineering")
                second.append(["Company", "Engineering Count", "Engineering Index"])
                second.append(["Other Corp", 20, 108])
                if with_chart:
                    chart = BarChart()
                    chart.add_data(Reference(first, min_col=2, max_col=3,
                                             min_row=1, max_row=2), titles_from_data=True)
                    chart.set_categories(Reference(first, min_col=1, min_row=2, max_row=2))
                    overview = book.create_chartsheet("Overview", 0)
                    overview.add_chart(chart)
                book.save(workbook)
            finally:
                book.close()

            result = subprocess.run(
                [sys.executable, str(TOOL), str(workbook), "--output", str(output)],
                cwd=root, capture_output=True, text=True,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            data = json.loads(output.read_text(encoding="utf-8"))
            self.assertEqual(data["companies"], [
                {"company": "Acme Corp", "city": "", "categories": {
                    "all_employees": {"count": 12, "index": 105.0}}},
                {"company": "Other Corp", "city": "", "categories": {
                    "engineering": {"count": 20, "index": 108.0}}},
            ])

    def test_chart_only_tab_does_not_abort_salary_conversion(self):
        self.convert(with_chart=True)

    def test_all_data_worksheets_are_preserved_without_chart_tabs(self):
        self.convert(with_chart=False)


if __name__ == "__main__":
    unittest.main()
