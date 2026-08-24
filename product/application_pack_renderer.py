"""Deterministic downstream renderer for a confirmed immutable Application Pack.

This module never re-queries the Evidence Profile, Job Fit, Application
Intelligence, or review state. It consumes exactly the ``cv_content``,
``cover_letter_content``, and ``job`` fields already embedded in a pack
payload and turns approved units into presentation documents. It never
invents application substance: it renders exactly the units the pack
contains, in the order the pack contains them, and nothing else.

The current Application Pack contract (``application-pack.v0``) does not
embed candidate identity/contact details or structured employer/role/date/
education/certification data -- those live in the Evidence Profile snapshot
and are referenced from the pack only by artifact ID. Rendering a
conventional full CV with a name/contact header and a structured experience
section requires that data to be captured into the pack itself at
confirmation time; this renderer intentionally does not reach upstream to
fill that gap, since doing so would make a rendered document derivable from
something other than the exact immutable pack it claims to represent. This
is a documented limitation of the current pack contract, not a rendering
choice.
"""
from __future__ import annotations

import hashlib
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

_FROZEN_ZIP_TIMESTAMP = (1980, 1, 1, 0, 0, 0)

RENDERER_VERSION = "application-pack-renderer.v1"

_CV_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_MAX_FILENAME_STEM_LENGTH = 120


class RendererError(RuntimeError):
    pass


@dataclass(frozen=True)
class RenderedFile:
    kind: str
    filename: str
    content: bytes
    mime_type: str
    content_hash: str


@dataclass(frozen=True)
class RenderedApplicationPack:
    source_pack_id: str
    renderer_version: str
    files: tuple[RenderedFile, ...]

    def file(self, kind: str) -> RenderedFile:
        for candidate in self.files:
            if candidate.kind == kind:
                return candidate
        raise RendererError(f"rendered application pack has no file of kind {kind!r}")


def _normalized_text(value: Any) -> str:
    return " ".join(unicodedata.normalize("NFKC", str(value or "")).split())


def _sanitize_filename_component(value: str, *, fallback: str) -> str:
    normalized = _normalized_text(value)
    stripped = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", normalized)
    collapsed = re.sub(r"[ \t]+", "_", stripped).strip("_.")
    return collapsed or fallback


def _content_hash(content: bytes) -> str:
    return f"sha256:{hashlib.sha256(content).hexdigest()}"


def _freeze_docx_bytes(content: bytes) -> bytes:
    """Rewrite a saved DOCX so its bytes depend only on document content.

    ``python-docx`` (via ``zipfile.ZipFile.writestr``) stamps every ZIP entry
    with the wall-clock time at save, so otherwise-identical documents saved
    a second apart produce different bytes. Rewriting every entry with a
    fixed timestamp removes that source of nondeterminism while leaving the
    OOXML content, entry order, and compression untouched, so ``content_hash``
    is a real raw-byte SHA-256 that is stable for identical renderer input.
    """

    source = zipfile.ZipFile(BytesIO(content))
    frozen_buffer = BytesIO()
    with zipfile.ZipFile(frozen_buffer, "w", zipfile.ZIP_DEFLATED) as frozen:
        for item in source.infolist():
            frozen_info = zipfile.ZipInfo(item.filename, date_time=_FROZEN_ZIP_TIMESTAMP)
            frozen_info.compress_type = item.compress_type
            frozen_info.external_attr = item.external_attr
            frozen.writestr(frozen_info, source.read(item.filename))
    return frozen_buffer.getvalue()


def _units_by_type(units: list[dict[str, Any]], unit_type: str) -> list[dict[str, Any]]:
    return [unit for unit in units if unit.get("unit_type") == unit_type]


def _build_filename_stem(pack: dict[str, Any]) -> str:
    job = pack.get("job") or {}
    company = _sanitize_filename_component(job.get("company", ""), fallback="Company")
    title = _sanitize_filename_component(job.get("title", ""), fallback="Role")
    stem = f"{company}_{title}"
    if len(stem) > _MAX_FILENAME_STEM_LENGTH:
        stem = stem[:_MAX_FILENAME_STEM_LENGTH].rstrip("_")
    return stem or "Application"


def _set_base_style(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for section in document.sections:
        section.top_margin = section.bottom_margin = Pt(54)
        section.left_margin = section.right_margin = Pt(54)


def _add_heading(document: Document, text: str) -> None:
    heading = document.add_heading(text, level=1)
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(6)


def render_cv_document(pack: dict[str, Any]) -> bytes:
    """Render the approved CV content of ``pack`` into a DOCX document.

    Only ``cv_summary_line`` and ``cv_bullet`` units already present in
    ``pack["cv_content"]`` are rendered, in pack order. No other section is
    added: the current pack contract carries no candidate name, contact
    details, or structured employer/role/date/education/certification data,
    so no such section can be faithfully rendered without inventing content.
    """

    cv_content = pack.get("cv_content") or []
    summary_lines = _units_by_type(cv_content, "cv_summary_line")
    bullets = _units_by_type(cv_content, "cv_bullet")

    document = Document()
    _set_base_style(document)

    job = pack.get("job") or {}
    title_text = " - ".join(
        part for part in (job.get("title"), job.get("company")) if part
    )
    if title_text:
        title_paragraph = document.add_paragraph(title_text)
        title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        title_run = title_paragraph.runs[0]
        title_run.bold = True
        title_run.font.size = Pt(14)

    if summary_lines:
        _add_heading(document, "Professional Summary")
        for unit in summary_lines:
            document.add_paragraph(unit.get("text", ""))

    if bullets:
        _add_heading(document, "Key Experience & Skills")
        for unit in bullets:
            document.add_paragraph(unit.get("text", ""), style="List Bullet")

    buffer = BytesIO()
    document.save(buffer)
    return _freeze_docx_bytes(buffer.getvalue())


def render_cover_letter_document(pack: dict[str, Any]) -> bytes:
    """Render the approved cover-letter content of ``pack`` into a DOCX document.

    Only ``cover_letter_paragraph`` units already present in
    ``pack["cover_letter_content"]`` are rendered, in pack order, as the body
    of a plain business-letter layout. No recipient name, address,
    salutation, or candidate contact details are invented: the pack contract
    carries no such fields today.
    """

    cover_letter_content = pack.get("cover_letter_content") or []
    paragraphs = _units_by_type(cover_letter_content, "cover_letter_paragraph")

    document = Document()
    _set_base_style(document)

    job = pack.get("job") or {}
    subject_parts = [part for part in (job.get("title"), job.get("company")) if part]
    if subject_parts:
        subject = document.add_paragraph(f"Re: {' - '.join(subject_parts)}")
        subject.runs[0].bold = True

    for unit in paragraphs:
        document.add_paragraph(unit.get("text", ""))

    buffer = BytesIO()
    document.save(buffer)
    return _freeze_docx_bytes(buffer.getvalue())


def render_application_pack(
    pack: dict[str, Any], *, source_pack_id: str
) -> RenderedApplicationPack:
    """Render both documents for ``pack``, an exact immutable pack payload.

    ``source_pack_id`` must be the artifact ID of the exact immutable
    Application Pack this payload was read from, so every rendered file is
    traceable to one specific pack version. This function performs no
    database access and no re-derivation of upstream artifacts: it only
    reads fields already present in ``pack``.

    Rendering the same pack content with the same renderer version produces
    byte-identical DOCX output, so each file's ``content_hash`` is a plain
    SHA-256 of its exact bytes: see ``_freeze_docx_bytes`` for how the
    otherwise wall-clock-dependent ZIP timestamps are removed.
    """

    if not isinstance(pack, dict):
        raise RendererError("pack must be a dict")
    if not source_pack_id:
        raise RendererError("source_pack_id is required for traceability")

    stem = _build_filename_stem(pack)

    cv_bytes = render_cv_document(pack)
    cover_letter_bytes = render_cover_letter_document(pack)

    files = (
        RenderedFile(
            kind="cv",
            filename=f"{stem}_CV.docx",
            content=cv_bytes,
            mime_type=_CV_DOCX_MIME,
            content_hash=_content_hash(cv_bytes),
        ),
        RenderedFile(
            kind="cover_letter",
            filename=f"{stem}_Cover_Letter.docx",
            content=cover_letter_bytes,
            mime_type=_CV_DOCX_MIME,
            content_hash=_content_hash(cover_letter_bytes),
        ),
    )
    return RenderedApplicationPack(
        source_pack_id=source_pack_id, renderer_version=RENDERER_VERSION, files=files,
    )
